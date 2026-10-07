#!/usr/bin/env python3
"""Generate .docx versions of the paperwork pages (no banners, no tracked changes, no comments)."""

import html as ihtml
import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from build import PAGES  # noqa
import pages_paperwork  # noqa: F401 (registers pages)
from docx import Document
from docx.shared import Pt, RGBColor

INK = RGBColor(0x10, 0x18, 0x28)
GOLD = RGBColor(0xB9, 0x8A, 0x2F)
MUTED = RGBColor(0x47, 0x54, 0x67)


def clean_text(s):
    s = re.sub(r"<[^>]+>", "", s)
    return ihtml.unescape(s).strip()


def add_runs(para, html):
    # minimal inline markup: <strong>, <em>, <code>, <a>
    tok = re.split(r"(<strong>.*?</strong>|<em>.*?</em>|<code.*?>.*?</code>|<a .*?>.*?</a>)", html)
    for t in tok:
        if not t:
            continue
        m = re.match(r"<(strong|em|code|a)(.*?)>(.*)</\1>$", t, re.S)
        if not m:
            para.add_run(ihtml.unescape(t))
            continue
        tag, attrs, inner = m.groups()
        text = ihtml.unescape(re.sub(r"<[^>]+>", "", inner))
        run = para.add_run(text)
        if tag == "strong":
            run.bold = True
        elif tag == "em":
            run.italic = True
        elif tag == "code":
            run.font.name = "Consolas"
            run.font.size = Pt(10)
        elif tag == "a":
            href = re.search(r'href="([^"]+)"', attrs)
            if href:
                text = f"{text} ({href.group(1)})"
                run.text = text


def body_to_docx(doc, body):
    # strip section/wrap/prose scaffolding divs but keep their inner content
    chunks = re.split(
        r'(<h1.*?</h1>|<h2>.*?</h2>|<p.*?</p>|<ul>.*?</ul>|<div class="note.*?</div>)',
        body,
        flags=re.S,
    )
    for ch in chunks:
        if not ch or not ch.strip():
            continue
        if ch.startswith("<h1"):
            h = doc.add_heading(clean_text(ch), level=1)
            for r in h.runs:
                r.font.color.rgb = INK
        elif ch.startswith("<h2"):
            h = doc.add_heading(clean_text(ch), level=2)
            for r in h.runs:
                r.font.color.rgb = INK
        elif ch.startswith('<p class="sub"'):
            p = doc.add_paragraph()
            add_runs(p, re.sub(r"^<p[^>]*>|</p>$", "", ch))
            for r in p.runs:
                r.font.size = Pt(12)
                r.font.color.rgb = MUTED
        elif ch.startswith("<p"):
            p = doc.add_paragraph()
            add_runs(p, re.sub(r"^<p[^>]*>|</p>$", "", ch))
        elif ch.startswith("<ul"):
            for li in re.findall(r"<li>(.*?)</li>", ch, re.S):
                p = doc.add_paragraph(style="List Bullet")
                add_runs(p, li)
        elif 'class="note' in ch:
            inner = re.sub(r"^<div[^>]*>|</div>$", "", ch, flags=re.S)
            p = doc.add_paragraph()
            add_runs(p, inner)
            for r in p.runs:
                r.font.color.rgb = GOLD
                r.bold = True


OUTDIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "paperwork-docx")
os.makedirs(OUTDIR, exist_ok=True)
made = []
for path, title, _desc, body, _ in PAGES:
    if not path.startswith("/paperwork/") or path == "/paperwork/":
        continue
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(11)
    t = doc.add_paragraph()
    r = t.add_run(title)
    r.bold = True
    r.font.size = Pt(20)
    r.font.color.rgb = INK
    s = doc.add_paragraph()
    r = s.add_run("AxiomCode \u2014 Last updated 26 September 2026")
    r.font.size = Pt(10)
    r.font.color.rgb = MUTED
    body_to_docx(doc, body)
    # hygiene: no tracked changes / comments / banners by construction; verify no DRAFT text
    full = "\n".join(p.text for p in doc.paragraphs)
    assert "DRAFT" not in full.upper() or "draft" in title.lower(), f"banner text in {title}"
    assert len(doc.paragraphs) > 3, f"empty doc {title}"
    fname = path.strip("/").replace("/", "-") + ".docx"
    doc.save(os.path.join(OUTDIR, fname))
    made.append(fname)
print("wrote:", ", ".join(made))
